"""Document pipeline (ticket C6): store -> extract -> chunk.

Source-agnostic by design: Etimad serves TSD/BOQ files only to logged-in
accounts (verified 2026-09-03 — the visitor attachments component is an empty
shell), so bytes arrive from whichever fetch lane exists (supplier-account
fetcher, extension relay per M1-7, or manual upload). Everything downstream
of `ingest_file` is identical regardless of source.

Storage: MinIO/S3 keyed by content hash (dedup); metadata + extracted text in
Postgres (`documents`, `doc_chunks`). Embeddings are left NULL — filled by a
separate embedder job once the model/provider decision (C6 decision record)
is made. Virus scanning is a hook: `scan_status` stays 'pending' until a
ClamAV worker flips it; serving to users must filter on scan_status='clean'.
"""
from __future__ import annotations

import hashlib
import io
import json
import logging
import re

import asyncpg

log = logging.getLogger("thaqip.documents")

BUCKET = "thaqip-docs"
CHUNK_CHARS = 2400
CHUNK_OVERLAP = 300

_EXT_MIME = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".txt": "text/plain",
}


def sniff_mime(file_name: str) -> str:
    for ext, mime in _EXT_MIME.items():
        if file_name.lower().endswith(ext):
            return mime
    return "application/octet-stream"


# ---------------------------------------------------------------- extraction

def extract_text(file_name: str, data: bytes) -> str:
    """Best-effort text extraction. Returns '' when nothing extractable."""
    name = file_name.lower()
    try:
        if name.endswith(".pdf"):
            return _extract_pdf(data)
        if name.endswith(".docx"):
            return _extract_docx(data)
        if name.endswith(".xlsx"):
            return _extract_xlsx(data)
        if name.endswith(".txt"):
            return data.decode("utf-8", errors="replace")
    except Exception as exc:  # noqa: BLE001 — extraction must never kill ingestion
        log.warning("extraction failed for %s: %r", file_name, exc)
    return ""


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"# {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def chunk_text(text: str) -> list[str]:
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    chunks, start = [], 0
    while start < len(text):
        end = min(len(text), start + CHUNK_CHARS)
        if end < len(text):  # prefer breaking on a newline near the boundary
            nl = text.rfind("\n", start + CHUNK_CHARS // 2, end)
            if nl != -1:
                end = nl
        chunks.append(text[start:end].strip())
        start = max(end - CHUNK_OVERLAP, start + 1)
        if end == len(text):
            break
    return [c for c in chunks if c]


# ------------------------------------------------------------------- storage

def make_minio(endpoint: str, access_key: str, secret_key: str, *, secure: bool = False):
    from minio import Minio

    client = Minio(endpoint, access_key=access_key, secret_key=secret_key, secure=secure)
    if not client.bucket_exists(BUCKET):
        client.make_bucket(BUCKET)
    return client


async def ingest_file(
    pool: asyncpg.Pool,
    minio,
    *,
    tender_pk: int,
    kind: str,          # tsd | boq | attachment | clarification | other
    file_name: str,
    data: bytes,
) -> int | None:
    """Store one file end-to-end. Returns documents.id, or None if duplicate."""
    sha = hashlib.sha256(data).hexdigest()
    storage_key = f"{kind}/{sha[:2]}/{sha}/{file_name}"
    mime = sniff_mime(file_name)

    existing = await pool.fetchval(
        "SELECT id FROM documents WHERE tender_id = $1 AND sha256 = $2", tender_pk, sha
    )
    if existing:
        return None

    minio.put_object(BUCKET, storage_key, io.BytesIO(data), len(data), content_type=mime)

    text = extract_text(file_name, data)
    chunks = chunk_text(text)

    async with pool.acquire() as conn, conn.transaction():
        doc_id = await conn.fetchval(
            """INSERT INTO documents (tender_id, kind, file_name, mime_type, size_bytes,
                                          sha256, storage_key, text_extracted)
                   VALUES ($1, $2, $3, $4, $5, $6, $7, $8) RETURNING id""",
            tender_pk, kind, file_name, mime, len(data), sha, storage_key, bool(chunks),
        )
        for i, chunk in enumerate(chunks):
            await conn.execute(
                "INSERT INTO doc_chunks (document_id, chunk_no, content) VALUES ($1, $2, $3)",
                doc_id, i, chunk,
            )
        await conn.execute(
            """INSERT INTO ingest_events (event_type, entity_type, entity_id, data)
                   VALUES ('document.stored', 'document', $1, $2::jsonb)""",
            doc_id,
            json.dumps({"tender_id": tender_pk, "kind": kind, "file_name": file_name,
                        "chunks": len(chunks)}, ensure_ascii=False),
        )
    log.info("stored %s (%s, %d bytes, %d chunks) for tender %d",
             file_name, kind, len(data), len(chunks), tender_pk)
    return doc_id
